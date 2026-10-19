#!/usr/bin/env python3
"""
gen_budget.py -- Genera budget/budget.xlsx con formato formal a partir de budget/budget.json.
"""
import json
import pathlib
import sys
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

BUDGET_JSON = pathlib.Path("budget/budget.json")
BUDGET_XLSX = pathlib.Path("budget/budget.xlsx")
BRIEF_JSON = pathlib.Path("brief.json")

def main():
    if not BUDGET_JSON.exists():
        print("ERROR: budget/budget.json no existe.")
        sys.exit(1)

    items = json.loads(BUDGET_JSON.read_text(encoding="utf-8"))
    brief = json.loads(BRIEF_JSON.read_text(encoding="utf-8")) if BRIEF_JSON.exists() else {}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presupuesto Puesta en Marcha"
    ws.views.sheetView[0].showGridLines = True

    # Estilos
    font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="595959")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    font_bold = Font(name="Calibri", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Titulos
    project_name = brief.get("project_name", "PROYECTO").upper()
    team_name = brief.get("team_name", "EQUIPO")
    ws.merge_cells("A1:H1")
    ws["A1"] = f"PRESUPUESTO DE PUESTA EN MARCHA (PILOTO INICIAL) -- {project_name}"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_left

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Equipo: {team_name} | Moneda base: Dolares Estadounidenses (USD) | Metodologia: Costos mayoristas sustentados"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = align_left

    # Encabezados de tabla
    headers = ["#", "Elemento / Insumo", "Descripcion Tecnica", "Unidades", "Costo Unit. (USD)", "Subtotal (USD)", "Fuente / Enlace", "Fecha"]
    row_idx = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    total_general = 0.0
    row_start = 5

    for i, it in enumerate(items, 1):
        curr_row = row_start + i - 1
        units = float(it.get("units", 1))
        unit_cost = float(it.get("unit_cost_usd", 0.0))
        subtotal = round(units * unit_cost, 2)
        total_general += subtotal

        ws.cell(row=curr_row, column=1, value=i).alignment = align_center
        ws.cell(row=curr_row, column=2, value=it.get("element", "")).alignment = align_left
        ws.cell(row=curr_row, column=3, value=it.get("description", "")).alignment = align_left
        ws.cell(row=curr_row, column=4, value=units).alignment = align_center
        c_unit = ws.cell(row=curr_row, column=5, value=unit_cost)
        c_unit.alignment = align_right
        c_unit.number_format = "$#,##0.00"
        
        c_sub = ws.cell(row=curr_row, column=6, value=subtotal)
        c_sub.alignment = align_right
        c_sub.number_format = "$#,##0.00"
        
        ws.cell(row=curr_row, column=7, value=it.get("url", "")).alignment = align_left
        ws.cell(row=curr_row, column=8, value=it.get("date", "")).alignment = align_center

        for col in range(1, 9):
            ws.cell(row=curr_row, column=col).border = thin_border

    # Fila de Total
    total_row = row_start + len(items)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    ws.cell(row=total_row, column=1, value="TOTAL PROYECTADO DE PUESTA EN MARCHA (USD)").alignment = align_right
    ws.cell(row=total_row, column=1).font = font_bold
    
    cell_total = ws.cell(row=total_row, column=6, value=total_general)
    cell_total.font = font_bold
    cell_total.alignment = align_right
    cell_total.number_format = "$#,##0.00"
    cell_total.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    for col in range(1, 9):
        ws.cell(row=total_row, column=col).border = thin_border

    # Autoajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    wb.save(BUDGET_XLSX)
    print(f"OK: {BUDGET_XLSX} generado exitosamente.")
    print(f"Total Presupuestado: ${total_general:,.2f} USD")

    if total_general < 4000:
        print(f"WARN: Presupuesto (${total_general:,.2f}) inferior al rango optimo de viabilidad ($4,000 - $6,000 USD).")
    elif total_general > 6000:
        print(f"WARN: Presupuesto (${total_general:,.2f}) excede el techo recomendado ($6,000 USD).")
    else:
        print(f"VALIDACION: Presupuesto dentro del rango ideal de victoria ($4,000 - $6,000 USD).")

if __name__ == "__main__":
    main()
